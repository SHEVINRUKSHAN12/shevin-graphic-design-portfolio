from pathlib import Path
from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

ROOT = Path(__file__).parent
ASSETS = ROOT / "assets"
OUT = ROOT / "Shevin_Rathnasinghe_Graphic_Design_Portfolio"

items = [
    ("A Plus Construction", "Social media campaign · Poster series", "featured-design.png"),
    ("DJI Mic Mini", "Prime Camera Rent · Product promotion", "prime-camera-dji-mic.jpg"),
    ("Creator Combo", "Prime Camera Rent · Product promotion", "prime-camera-combo.jpg"),
    ("YouTube Channel", "St. Jude’s Media Crew · Social campaign", "st-judes-youtube.jpg"),
    ("Season 03 Final", "ASG Premier League · Match summary", "asgpremier-s03-summary.png"),
    ("Final Match", "ASG Premier League · Match artwork", "asgpremier-s03-final.png"),
    ("Player of the Match", "ASG Premier League · Sports graphic", "asgpremier-s03-player.png"),
    ("Match 05", "ASG Premier League · Match artwork", "asgpremier-match-05.png"),
    ("Match 06", "ASG Premier League · Match artwork", "asgpremier-match-06.png"),
    ("Match 04", "ASG Premier League · Match artwork", "asgpremier-match-04.png"),
    ("Match 03", "ASG Premier League · Match artwork", "asgpremier-match-03.png"),
    ("Match 02", "ASG Premier League · Match artwork", "asgpremier-match-02.png"),
    ("Match 01", "ASG Premier League · Match artwork", "asgpremier-match-01.png"),
    ("Ceylon Strikers", "ASG Premier League · Team reveal", "asgpremier-team-01.png"),
    ("Royal Gladiators", "ASG Premier League · Team reveal", "asgpremier-team-02.png"),
    ("Kingsford XI", "ASG Premier League · Team reveal", "asgpremier-team-03.png"),
    ("Phantom Sports CC", "ASG Premier League · Team reveal", "asgpremier-team-04.png"),
    ("Liberty CC", "ASG Premier League · Team reveal", "asgpremier-team-05.png"),
    ("Lion Kings", "ASG Premier League · Team reveal", "asgpremier-team-06.png"),
    ("Fury Falcons", "ASG Premier League · Team reveal", "asgpremier-team-07.png"),
    ("Player Welcome", "ASG Premier League · Team reveal", "asgpremier-team-08.png"),
    ("Garcinia Herb", "Ranwan Lanka · Product campaign", "ranwan-garcinia-01.jpg"),
    ("Women’s Day", "Ranwan Lanka · Campaign artwork", "ranwan-womens-day.jpg"),
    ("Product Promotion", "Ranwan Lanka · Social creative", "ranwan-garcinia-02.jpg"),
    ("Sinhala Campaign", "Ranwan Lanka · Promotional design", "ranwan-sinhala-promo.jpg"),
    ("Aurudu Discount", "Ranwan Lanka · Seasonal campaign", "ranwan-aurudu.jpg"),
    ("Fabric Conditioner", "Sanirub · Product campaign", "sanirub-conditioner.jpg"),
    ("Surface Cleaner", "Sanirub · Product campaign", "sanirub-floor-cleaner.jpg"),
    ("Air Freshener", "Sanirub · Product campaign", "sanirub-air-freshener.jpg"),
    ("Christmas Sale", "Sanirub · Seasonal campaign", "sanirub-christmas.jpg"),
]

for _, _, filename in items:
    if not (ASSETS / filename).exists():
        raise FileNotFoundError(filename)

def add_doc_image(cell, path, max_width=4.85, max_height=5.55):
    with Image.open(path) as im:
        width, height = im.size
    scale = min(max_width / width, max_height / height)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width * scale), height=Inches(height * scale))

def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is None:
        zoom = OxmlElement("w:zoom")
        doc.settings.element.append(zoom)
    zoom.set(qn("w:percent"), "100")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Inches(1.45)
    run = title.add_run("SHEVIN RATHNASINGHE")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(32)
    run.font.color.rgb = RGBColor(177, 31, 75)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("GRAPHIC DESIGN PORTFOLIO").bold = True
    subtitle.runs[0].font.size = Pt(16)
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.space_before = Inches(0.35)
    intro.add_run("Selected posters, social media designs, campaign artwork and visual concepts.").italic = True
    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_before = Inches(2.1)
    contact.add_run("shevinrukshan12@gmail.com").bold = True
    doc.add_page_break()

    for index in range(0, len(items), 2):
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width = Inches(5.1)
        table.columns[1].width = Inches(5.1)
        for col, item in enumerate(items[index:index + 2]):
            cell = table.cell(0, col)
            cell.width = Inches(5.1)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            add_doc_image(cell, ASSETS / item[2])
            caption = cell.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.space_before = Pt(5)
            caption.add_run(item[0]).bold = True
            detail = cell.add_paragraph(item[1])
            detail.alignment = WD_ALIGN_PARAGRAPH.CENTER
            detail.runs[0].font.size = Pt(8)
            detail.runs[0].font.color.rgb = RGBColor(92, 92, 92)
        if index + 2 < len(items):
            doc.add_page_break()
    doc.core_properties.title = "Shevin Rathnasinghe — Graphic Design Portfolio"
    doc.core_properties.author = "Shevin Rathnasinghe"
    doc.save(str(OUT.with_suffix(".docx")))

def fit_image(path, max_w, max_h):
    with Image.open(path) as im:
        width, height = im.size
    scale = min(max_w / width, max_h / height)
    return width * scale, height * scale

def build_pdf():
    page_w, page_h = landscape(A4)
    pdf = canvas.Canvas(str(OUT.with_suffix(".pdf")), pagesize=(page_w, page_h))
    pdf.setTitle("Shevin Rathnasinghe — Graphic Design Portfolio")
    pdf.setAuthor("Shevin Rathnasinghe")
    pdf.setFillColor(colors.HexColor("#b11f4b"))
    pdf.setFont("Helvetica-Bold", 28)
    pdf.drawCentredString(page_w / 2, page_h * 0.66, "SHEVIN RATHNASINGHE")
    pdf.setFillColor(colors.HexColor("#242424"))
    pdf.setFont("Helvetica-Bold", 15)
    pdf.drawCentredString(page_w / 2, page_h * 0.58, "GRAPHIC DESIGN PORTFOLIO")
    pdf.setFont("Helvetica-Oblique", 11)
    pdf.drawCentredString(page_w / 2, page_h * 0.51, "Selected posters, social media designs, campaign artwork and visual concepts.")
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawCentredString(page_w / 2, page_h * 0.28, "shevinrukshan12@gmail.com")
    pdf.showPage()

    for index in range(0, len(items), 2):
        for col, item in enumerate(items[index:index + 2]):
            x = 30 + col * (page_w / 2)
            box_w = page_w / 2 - 55
            box_h = page_h - 86
            image_w, image_h = fit_image(ASSETS / item[2], box_w, box_h - 38)
            image_x = x + (box_w - image_w) / 2
            image_y = 51 + 25 + (box_h - 38 - image_h) / 2
            pdf.drawImage(ImageReader(str(ASSETS / item[2])), image_x, image_y, image_w, image_h, preserveAspectRatio=True, mask="auto")
            pdf.setFillColor(colors.HexColor("#242424"))
            pdf.setFont("Helvetica-Bold", 10)
            pdf.drawCentredString(x + box_w / 2, 38, item[0])
            pdf.setFillColor(colors.HexColor("#5c5c5c"))
            pdf.setFont("Helvetica", 7.5)
            pdf.drawCentredString(x + box_w / 2, 26, item[1])
        pdf.showPage()
    pdf.save()

build_docx()
build_pdf()
print(f"Created {OUT.with_suffix('.docx').name} and {OUT.with_suffix('.pdf').name}")
